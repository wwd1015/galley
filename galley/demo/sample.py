"""A sample legacy whitepaper in Word format, for the demo and for tests.

It has what a real one has: styled callouts, a picture, a table, captions,
cross-references, citations and (optionally) a native Word chart with its data.
"""

from __future__ import annotations

from itertools import pairwise
from pathlib import Path
from typing import Any

import docx
import pymupdf
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches

INTRO = (
    "This review covers deposit behaviour over the 30-day stress horizon. "
    "Retail balances fell by 4.5% while wholesale balances fell by 12.25%."
)
SENSITIVITY = (
    "Wholesale funding is more sensitive to rating actions than retail funding, "
    "and it reprices faster after a downgrade."
)
KEY_FINDING = "The buffer covers 118% of net outflows in every scenario."
SCOPE = "The review covers 3 portfolios with combined balances of $24,805m."
BULLETS = ["retail and small business deposits;", "wholesale and operational deposits."]
RESULTS = "Figure 1 shows the buffer and Table 1 lists the outflow rates."
FIGURE_CAPTION = "Liquidity buffer over the stress horizon"
TABLE_CAPTION = "Outflow rates by category"
TABLE = [
    ["Category", "Rate", "Balance"],
    ["Stable retail", "5%", "12,400"],
    ["Less stable retail", "10%", "6,150"],
    ["Operational", "25%", "6,255"],
]
CLOSING = "Rates were last calibrated in 2025."


def make_png(path: Path) -> Path:
    """A small chart-like image."""
    document = pymupdf.open()
    page = document.new_page(width=360, height=200)
    page.draw_line((30, 170), (340, 170), color=(0.4, 0.4, 0.4), width=1)
    page.draw_line((30, 170), (30, 20), color=(0.4, 0.4, 0.4), width=1)
    points = [(30, 150), (120, 110), (210, 120), (340, 40)]
    for start, end in pairwise(points):
        page.draw_line(start, end, color=(0.08, 0.2, 0.31), width=2)
    path.parent.mkdir(parents=True, exist_ok=True)
    page.get_pixmap(dpi=100).save(path)
    return path


CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
CHART_TYPE = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
CHART_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
CHART_CATEGORIES = ["Q1", "Q2", "Q3", "Q4"]
CHART_SERIES = {"Retail": ["101", "104", "99", "107"], "Wholesale": ["96", "92", "88", "91"]}
CHART_CAPTION = "Deposit balances by quarter"


def chart_xml(title: str, categories: list[str], series: dict[str, list[str]]) -> bytes:
    def cache(tag: str, kind: str, values: list[str]) -> str:
        points = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(values))
        return (
            f"<c:{tag}><c:{kind}Ref><c:{kind}Cache>{points}</c:{kind}Cache></c:{kind}Ref></c:{tag}>"
        )

    body = "".join(
        f'<c:ser><c:idx val="{i}"/>{cache("tx", "str", [name])}'
        f"{cache('cat', 'str', categories)}{cache('val', 'num', values)}</c:ser>"
        for i, (name, values) in enumerate(series.items())
    )
    return (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<c:chartSpace xmlns:c="{CHART_NS}" '
        f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart>'
        f"<c:title><c:tx><c:rich><a:p><a:r><a:t>{title}</a:t></a:r></a:p></c:rich></c:tx></c:title>"
        f"<c:plotArea><c:lineChart>{body}</c:lineChart></c:plotArea></c:chart></c:chartSpace>"
    ).encode()


def add_chart(document: Any, title: str) -> None:
    """Embed a native Word chart (data cached in the chart part, as Word writes it)."""
    part = Part(
        PackURI("/word/charts/chart1.xml"),
        CHART_TYPE,
        chart_xml(title, CHART_CATEGORIES, CHART_SERIES),
        document.part.package,
    )
    rel_id = document.part.relate_to(part, CHART_REL)
    drawing = parse_xml(
        f"<w:drawing {nsdecls('w', 'wp', 'a', 'r')}>"
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        '<wp:extent cx="4572000" cy="2743200"/><wp:docPr id="50" name="Chart 1"/>'
        f'<a:graphic><a:graphicData uri="{CHART_NS}">'
        f'<c:chart xmlns:c="{CHART_NS}" r:id="{rel_id}"/>'
        "</a:graphicData></a:graphic></wp:inline></w:drawing>"
    )
    document.add_paragraph().add_run()._r.append(drawing)


def make_docx(
    path: Path,
    image: Path,
    *,
    references: bool = False,
    chart: bool = False,
    extras: bool = False,
) -> Path:
    """The legacy Word whitepaper the fixtures pretend to migrate."""
    document = docx.Document()
    document.styles.add_style("Key Finding", WD_STYLE_TYPE.PARAGRAPH)
    document.styles.add_style("Pull Quote", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Deposit Outflow Review", style="Title")
    document.add_heading("Introduction", level=1)
    document.add_paragraph(INTRO)
    document.add_paragraph(SENSITIVITY)
    document.add_paragraph(KEY_FINDING, style="Key Finding")
    document.add_heading("Scope", level=2)
    document.add_paragraph(SCOPE)
    for bullet in BULLETS:
        document.add_paragraph(bullet, style="List Bullet")
    document.add_heading("Results", level=1)
    document.add_paragraph(RESULTS)
    document.add_picture(str(image), width=Inches(4))
    document.add_paragraph(f"Figure 1. {FIGURE_CAPTION}", style="Caption")
    document.add_paragraph(f"Table 1. {TABLE_CAPTION}", style="Caption")
    table = document.add_table(rows=len(TABLE), cols=len(TABLE[0]))
    for row, values in zip(table.rows, TABLE, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph(CLOSING)
    if chart:
        document.add_paragraph("Figure 2 compares retail and wholesale balances.")
        add_chart(document, "Balances")
        document.add_paragraph(f"Figure 2. {CHART_CAPTION}", style="Caption")
    if extras:
        document.add_paragraph("Funding costs rose over the year.", style="Pull Quote")
    if references:
        document.add_paragraph(
            "The framework follows the Basel standard (Basel Committee 2013) and the "
            "classic model of runs (Diamond and Dybvig 1983); see also (Unknown 1999)."
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path
