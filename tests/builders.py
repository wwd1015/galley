"""Builders for test documents: a Word source and the paper it converts to."""

from __future__ import annotations

from itertools import pairwise
from pathlib import Path

import docx
import pymupdf
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches

from galley import manifest, scaffold

INTRO = (
    "This review covers deposit behaviour over the 30-day stress horizon. "
    "Retail balances fell by 4.5% while wholesale balances fell by 12.25%."
)
SENSITIVITY = (
    "Wholesale funding is more sensitive to rating actions than retail funding, "
    "and it reprices faster after a downgrade."
)
KEY_FINDING = "The buffer covers 118% of net outflows in every scenario."
SCOPE = "The review covers 3 portfolios with combined balances of $24,805m."
BULLETS = ["retail and small business deposits;", "wholesale and operational deposits."]
RESULTS = "Figure 1 shows the buffer and Table 1 lists the outflow rates."
FIGURE_CAPTION = "Liquidity buffer over the stress horizon"
TABLE_CAPTION = "Outflow rates by category"
TABLE = [
    ["Category", "Rate", "Balance"],
    ["Stable retail", "5%", "12,400"],
    ["Less stable retail", "10%", "6,150"],
    ["Operational", "25%", "6,255"],
]
CLOSING = "Rates were last calibrated in 2025."

PAPER_QMD = f"""\
---
title: "Deposit Outflow Review"
author:
  - Risk Analytics Team
date: "2026-10-05"
document-id: WP-TEST-002
classification: Internal
version: "0.3"
---

# Introduction

{INTRO}

{SENSITIVITY}

::: {{.keyfinding}}
{KEY_FINDING}
:::

## Scope

{SCOPE.replace("$", chr(92) + "$")}

- {BULLETS[0]}
- {BULLETS[1]}

# Results

@fig-buffer shows the buffer and @tbl-rates lists the outflow rates.

![{FIGURE_CAPTION}](figures/source/image1.png){{#fig-buffer width=70%}}

| Category           | Rate | Balance |
|:-------------------|-----:|--------:|
| Stable retail      |   5% |  12,400 |
| Less stable retail |  10% |   6,150 |
| Operational        |  25% |   6,255 |

: {TABLE_CAPTION} {{#tbl-rates}}

{CLOSING}
"""


def make_png(path: Path) -> Path:
    """A small chart-like image."""
    document = pymupdf.open()
    page = document.new_page(width=360, height=200)
    page.draw_line((30, 170), (340, 170), color=(0.4, 0.4, 0.4), width=1)
    page.draw_line((30, 170), (30, 20), color=(0.4, 0.4, 0.4), width=1)
    points = [(30, 150), (120, 110), (210, 120), (340, 40)]
    for start, end in pairwise(points):
        page.draw_line(start, end, color=(0.08, 0.2, 0.31), width=2)
    path.parent.mkdir(parents=True, exist_ok=True)
    page.get_pixmap(dpi=100).save(path)
    return path


def make_docx(path: Path, image: Path, *, references: bool = False) -> Path:
    """The legacy Word whitepaper the fixtures pretend to migrate."""
    document = docx.Document()
    document.styles.add_style("Key Finding", WD_STYLE_TYPE.PARAGRAPH)
    document.styles.add_style("Pull Quote", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Deposit Outflow Review", style="Title")
    document.add_heading("Introduction", level=1)
    document.add_paragraph(INTRO)
    document.add_paragraph(SENSITIVITY)
    document.add_paragraph(KEY_FINDING, style="Key Finding")
    document.add_heading("Scope", level=2)
    document.add_paragraph(SCOPE)
    for bullet in BULLETS:
        document.add_paragraph(bullet, style="List Bullet")
    document.add_heading("Results", level=1)
    document.add_paragraph(RESULTS)
    document.add_picture(str(image), width=Inches(4))
    document.add_paragraph(f"Figure 1. {FIGURE_CAPTION}", style="Caption")
    document.add_paragraph(f"Table 1. {TABLE_CAPTION}", style="Caption")
    table = document.add_table(rows=len(TABLE), cols=len(TABLE[0]))
    for row, values in zip(table.rows, TABLE, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph(CLOSING)
    if references:
        document.add_paragraph(
            "The framework follows the Basel standard (Basel Committee 2013) and the "
            "classic model of runs (Diamond and Dybvig 1983); see also (Unknown 1999)."
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def make_paper(paper_dir: Path, qmd: str = PAPER_QMD) -> Path:
    """A paper repo holding the hand-converted equivalent of :func:`make_docx`."""
    paper_dir.mkdir(parents=True, exist_ok=True)
    scaffold.vendor_extension(paper_dir)
    scaffold.write_project_files(paper_dir, paper_dir.name)
    manifest.save(paper_dir, [])
    make_png(paper_dir / "figures" / "source" / "image1.png")
    (paper_dir / "paper.qmd").write_text(qmd, encoding="utf-8")
    return paper_dir
